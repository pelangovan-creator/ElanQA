"""Generates the synthetic sample requirement documents used by tests and demos.
Run: python3 apps/api/scripts/make-fixtures.py  (needs python-docx, reportlab, openpyxl, pillow)"""
import csv, io, os
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw, ImageFont
import openpyxl

OUT = os.path.join(os.path.dirname(__file__), '..', 'test', 'fixtures')
os.makedirs(OUT, exist_ok=True)

INTRO = 'This document describes the requirements for the ShopEasy web store. It is intended for the QA and development teams.'
LOGIN = [
    'The system shall allow registered users to log in with email and password.',
    'The account must be locked after 3 consecutive failed login attempts.',
    'The login page should load fast on mobile devices.',
]
BULLETS = ['Password must be 8 to 16 characters long', 'Password must contain at least one digit', 'Social login is out of scope for this release']
TABLE = [['Field', 'Type', 'Validation'], ['Email', 'Text', 'Mandatory, valid email format'], ['Card number', 'Numeric', '16 digits, Luhn check']]
CHECKOUT = ['Orders above Rs. 5000 receive a 10% discount.', 'Payments are processed through the Razorpay API.', 'The delivery charge is TBD.']

def docx(path, v2=False):
    d = Document()
    d.add_heading('ShopEasy Software Requirements Specification', 0)
    d.add_heading('1 Introduction', 1)
    d.add_paragraph(INTRO)
    d.add_heading('2 Login', 1)
    login = list(LOGIN)
    if v2:
        login[1] = 'The account must be locked after 5 consecutive failed login attempts.'
        login.append('An OTP shall be sent by SMS for new devices.')
    d.add_paragraph(' '.join(login))
    for b in BULLETS if not v2 else BULLETS[:2]:
        d.add_paragraph(b, style='List Bullet')
    d.add_paragraph('Table 1: Payment form fields')
    t = d.add_table(rows=len(TABLE), cols=3)
    for r, row in enumerate(TABLE):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.add_heading('3 Checkout', 1)
    d.add_paragraph(' '.join(CHECKOUT))
    d.add_paragraph('Note: Prices are shown in INR.')
    d.add_heading('Appendix A Glossary', 1)
    d.add_paragraph('OTP means one-time password.')
    d.save(path)

def pdf(path):
    c = canvas.Canvas(path, pagesize=A4)
    y = 800
    for line in ['2 Login', *LOGIN, '', '- ' + BULLETS[0], '- ' + BULLETS[1]]:
        c.drawString(60, y, line); y -= 18
    c.showPage()
    y = 800
    for line in ['3 Checkout', *CHECKOUT]:
        c.drawString(60, y, line); y -= 18
    c.showPage()
    c.save()

def text_image(lines, size=(1400, 60 + 50 * 6)):
    img = Image.new('L', size, 255)
    dr = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 30)
    except OSError:
        font = ImageFont.load_default()
    y = 30
    for l in lines:
        dr.text((40, y), l, fill=0, font=font); y += 50
    return img

def scanned_pdf(path):
    img = text_image(['2 Login', *LOGIN])
    buf = io.BytesIO(); img.save(buf, format='PNG'); buf.seek(0)
    from reportlab.lib.utils import ImageReader
    c = canvas.Canvas(path, pagesize=(img.width, img.height))
    c.drawImage(ImageReader(buf), 0, 0, img.width, img.height)
    c.showPage(); c.save()

def xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'User Stories'
    ws.append(['ID', 'As a', 'I want to', 'So that'])
    ws.append(['US-1', 'shopper', 'save items to a wishlist', 'I can buy them later'])
    ws.append(['US-2', 'admin', 'export orders to CSV', 'finance can reconcile payments'])
    wb.save(path)

def md(path):
    with open(path, 'w') as f:
        f.write('# Search\n\nUsers can search products by name. Results should appear quickly, etc.\n\n'
                '| Filter | Values |\n|---|---|\n| Price | 0 - 100000 |\n| Rating | 1 to 5 stars |\n\n'
                '> Search history is stored for approximately 30 days.\n\n[^1]: Elastic search is assumed to be available.\n')

docx(os.path.join(OUT, 'srs-v1.docx'))
docx(os.path.join(OUT, 'srs-v2.docx'), v2=True)
pdf(os.path.join(OUT, 'login-checkout.pdf'))
scanned_pdf(os.path.join(OUT, 'scanned-login.pdf'))
text_image(['2 Login', *LOGIN]).save(os.path.join(OUT, 'login-screenshot.png'))
xlsx(os.path.join(OUT, 'user-stories.xlsx'))
md(os.path.join(OUT, 'search.md'))
with open(os.path.join(OUT, 'stories.csv'), 'w', newline='') as f:
    csv.writer(f).writerows([['ID', 'Story'], ['US-9', 'As a guest I want to track my order, so that I know the delivery date.']])
with open(os.path.join(OUT, 'email.txt'), 'w') as f:
    f.write('Hi team,\n\nPlease add a "Buy again" button on the order history page. It may need to support bulk orders.\n\nThanks,\nRavi\n')
print('fixtures written to', os.path.abspath(OUT))
